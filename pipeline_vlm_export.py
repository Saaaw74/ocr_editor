# -*- coding: utf-8 -*-


from __future__ import annotations

import os
import io
import re
import math
import time
import base64
import logging
import threading
from concurrent.futures import ThreadPoolExecutor
import requests
from html.parser import HTMLParser
import html

import fitz  # PyMuPDF
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

logger = logging.getLogger("pipeline_vlm_export")

# ----------------- КОНФИГУРАЦИЯ ИЗ .ENV -----------------
LIGHTON_API_KEY = os.getenv("LIGHTON_API_KEY", "")
LIGHTON_API_URL = os.getenv("LIGHTON_API_URL", "https://api.lighton.ai/api/v3/parse")
LIGHTON_TIMEOUT_SEC = int(os.getenv("LIGHTON_TIMEOUT_SEC", "120"))
OCR_RETRY_DELAY_SEC = int(os.getenv("OCR_RETRY_DELAY_SEC", "2"))

OPENROUTER_API_KEY = os.getenv("OPENROUTER_API_KEY", "")
OPENROUTER_API_URL = os.getenv("OPENROUTER_API_URL", "https://openrouter.ai/api/v1/chat/completions")
QWEN_MODEL = os.getenv("QWEN_MODEL", "qwen/qwen3.5-9b")
QWEN_FALLBACK_MODEL = os.getenv("QWEN_FALLBACK_MODEL", "qwen/qwen2.5-vl-72b-instruct")
QWEN_TIMEOUT_SEC = int(os.getenv("QWEN_TIMEOUT_SEC", "180"))

PAGE_DPI = int(os.getenv("PAGE_DPI", "200"))
PT_TO_CM = 2.54 / 72.0
DEFAULT_FONT_NAME = "Times New Roman"
DEFAULT_FONT_SIZE = 11.0
DEFAULT_MARGIN_CM = 1.2
PAGE_WIDTH_CM = 21.0
PAGE_HEIGHT_CM = 29.7


def _env_workers(*keys: str) -> int:
    """Читает *_MAX_WORKERS из окружения, значение по умолчанию строго 1.
    Возвращает max>=1 среди переданных ключей (архитектура проекта использует
    один пул для подготовки страниц экспорта)."""
    workers = 1
    for key in keys:
        try:
            workers = max(workers, int(os.getenv(key, "1") or "1"))
        except (TypeError, ValueError):
            continue
    return max(1, workers)



VLM_MAX_WORKERS = _env_workers("VLM_MAX_WORKERS")
DOCX_MAX_WORKERS = _env_workers("DOCX_MAX_WORKERS")

_ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

SYSTEM_PROMPT = (
    "Ты — универсальный экспертный модуль восстановления структуры документов для экспорта в Word.\n"
    "Тебе передан OCR-текст (Markdown) одной страницы и её изображение.\n"
    "Твоя задача — восстановить точную верстку страницы и вернуть строго валидный JSON по схеме:\n"
    "{\n"
    '  "elements": [\n'
    '    {"type": "heading", "level": 1, "text": "...", "bold": true, "font_size": 14, "align": "center|left|right", "runs": [{"text": "...", "bold": true}]},\n'
    '    {"type": "paragraph", "text": "...", "bold": false, "font_size": 10, "align": "left|center|right|justify", "runs": [{"text": "обычный ", "bold": false}, {"text": "жирный", "bold": true}]},\n'
    '    {"type": "table", "table_kind": "other|bank", "header_row": true|false, "bordered": true|false, "column_widths_cm": [9.0, 9.0], "rows": [["...", "..."]]}\n'
    "  ]\n"
    "}\n\n"
    "ПРАВИЛА ДЛЯ ТАБЛИЦ (bordered):\n"
    "- bordered=true, если у таблицы реально видны линии сетки/рамки ячеек.\n"
    "- bordered=false, если это параллельные колонки без линий (реквизиты, блок подписей с прочерками).\n\n"
    "УНИВЕРСАЛЬНЫЕ ПРАВИЛА ГЕОМЕТРИИ И ВЕРСТКИ:\n"
    "1. ВЫРАВНИВАНИЕ (align): Заголовки -> 'center', реквизиты справа -> 'right', текст -> 'left'|'justify'.\n"
    "2. МНОГОКОЛОНОЧНЫЕ БЛОКИ: Параллельные колонки (реквизиты, подписи) объединяй в table с table_kind='bank', bordered=false.\n"
    "3. СЕТОЧНЫЕ ТАБЛИЦЫ: Таблицы с линиями -> table_kind='other', header_row=true, bordered=true.\n"
    "4. ВЫДЕЛЕНИЯ И ШРИФТЫ: Сохраняй частичные жирные фрагменты в массиве 'runs'.\n"
    "5. ТОЧНОСТЬ ТЕКСТА: Не удаляй, не сокращай и не придумывай текст.\n"
    "6. ФОРМАТ ОТВЕТА: Верни ТОЛЬКО чистый JSON-объект без обёрток (```json) и комментариев."
)

# ----------------- СХЕМА ДАННЫХ -----------------
def make_paragraph(text: str, bold: bool = False, font_size: float = DEFAULT_FONT_SIZE,
                   align: str = "left", runs: list | None = None,
                   first_line_indent_cm: float = 0.0) -> dict:
    return {
        "type": "paragraph",
        "text": text or "",
        "bold": bool(bold),
        "font_size": font_size,
        "align": align if align in {"left", "center", "right", "justify"} else "left",
        "runs": runs,
        "first_line_indent_cm": first_line_indent_cm,
    }

def make_heading(text: str, level: int = 1, bold: bool = True, font_size: float = 16.0,
                 align: str = "left", runs: list | None = None) -> dict:
    return {
        "type": "heading",
        "level": max(1, min(6, int(level or 1))),
        "text": text or "",
        "bold": bool(bold),
        "font_size": font_size,
        "align": align if align in {"left", "center", "right", "justify"} else "left",
        "runs": runs,
        "first_line_indent_cm": 0.0,
    }

def make_table(rows: list, header_row: bool = True, table_kind: str = "other",
               column_widths_cm: list | None = None, bordered: bool = True) -> dict:
    clean_rows = [[str(cell) if cell is not None else "" for cell in row] for row in rows]
    return {
        "type": "table",
        "table_kind": table_kind if table_kind in {"products", "bank", "other"} else "other",
        "header_row": bool(header_row),
        "bordered": bool(bordered),
        "column_widths_cm": column_widths_cm,
        "rows": clean_rows,
    }

def normalize_elements(raw_elements) -> list[dict]:
    if not isinstance(raw_elements, list):
        return []
    result = []
    for raw in raw_elements:
        if not isinstance(raw, dict):
            continue
        etype = raw.get("type", "paragraph")
        if etype == "table":
            rows = raw.get("rows")
            if isinstance(rows, list) and rows:
                result.append(make_table(
                    rows=rows,
                    header_row=bool(raw.get("header_row", True)),
                    table_kind=raw.get("table_kind", "other"),
                    column_widths_cm=raw.get("column_widths_cm"),
                    bordered=bool(raw.get("bordered", True)),
                ))
        elif etype == "heading":
            result.append(make_heading(
                text=str(raw.get("text", "")),
                level=raw.get("level", 1),
                bold=raw.get("bold", True),
                font_size=raw.get("font_size", 16.0),
                align=raw.get("align", "left"),
                runs=raw.get("runs"),
            ))
        else:
            result.append(make_paragraph(
                text=str(raw.get("text", "")),
                bold=raw.get("bold", False),
                font_size=raw.get("font_size", DEFAULT_FONT_SIZE),
                align=raw.get("align", "left"),
                runs=raw.get("runs"),
                first_line_indent_cm=raw.get("first_line_indent_cm", 0.0),
            ))
    return result

# ----------------- КЛИЕНТ LIGHTON OCR -----------------
def _extract_lighton_markdown(payload) -> str | None:
    if isinstance(payload, str):
        return payload if payload.strip() else None
    if not isinstance(payload, dict):
        return None
    for key in ("markdown", "content", "text", "output", "result_text", "result", "parsed"):
        val = payload.get(key)
        if isinstance(val, str) and val.strip():
            return val
        if isinstance(val, dict):
            res = _extract_lighton_markdown(val)
            if res:
                return res
    for list_key in ("pages", "results", "data", "chunks"):
        val = payload.get(list_key)
        if isinstance(val, list) and val:
            parts = [s for item in val if (s := _extract_lighton_markdown(item))]
            if parts:
                return "\n\n---\n\n".join(parts)
    return None

def ocr_page_lighton(image_bytes: bytes, max_retries: int = 3) -> str:
    if not (LIGHTON_API_KEY and LIGHTON_API_URL):
        return ""
    headers = {"Authorization": f"Bearer {LIGHTON_API_KEY}"}
    for attempt_idx in range(max_retries):
        try:
            resp = requests.post(
                LIGHTON_API_URL,
                headers=headers,
                files={"file": ("page.png", image_bytes, "image/png")},
                timeout=LIGHTON_TIMEOUT_SEC,
            )
            if resp.status_code == 200:
                md = _extract_lighton_markdown(resp.json())
                if md:
                    return md
            if resp.status_code == 429:
                time.sleep(OCR_RETRY_DELAY_SEC * (attempt_idx + 1))
                continue
            b64 = base64.b64encode(image_bytes).decode("ascii")
            resp_json = requests.post(
                LIGHTON_API_URL,
                headers={"Authorization": f"Bearer {LIGHTON_API_KEY}", "Content-Type": "application/json"},
                json={"document": f"data:image/png;base64,{b64}"},
                timeout=LIGHTON_TIMEOUT_SEC,
            )
            if resp_json.status_code == 200:
                md = _extract_lighton_markdown(resp_json.json())
                if md:
                    return md
        except Exception as e:
            logger.warning(f"[LIGHTON] Попытка {attempt_idx+1} не удалась: {e}")
            time.sleep(1)
    return ""

# ----------------- КЛИЕНТ QWEN (OPENROUTER) -----------------
def refine_page_qwen(markdown_text: str, image_bytes: bytes | None = None) -> list[dict]:
    if not (OPENROUTER_API_KEY and OPENROUTER_API_URL and QWEN_MODEL):
        return []
    content = [{"type": "text", "text": markdown_text or "(пустая страница)"}]
    if image_bytes:
        b64 = base64.b64encode(image_bytes).decode("ascii")
        content.append({"type": "image_url", "image_url": {"url": f"data:image/png;base64,{b64}"}})

    headers = {"Authorization": f"Bearer {OPENROUTER_API_KEY}", "Content-Type": "application/json"}
    for model in filter(None, [QWEN_MODEL, QWEN_FALLBACK_MODEL]):
        try:
            payload = {
                "model": model,
                "messages": [{"role": "system", "content": SYSTEM_PROMPT}, {"role": "user", "content": content}],
                "temperature": 0.0,
                "response_format": {"type": "json_object"},
            }
            resp = requests.post(OPENROUTER_API_URL, headers=headers, json=payload, timeout=QWEN_TIMEOUT_SEC)
            if resp.status_code == 200:
                data = resp.json()["choices"][0]["message"]["content"]
                import json
                parsed = json.loads(data)
                return normalize_elements(parsed.get("elements"))
        except Exception as exc:
            logger.warning(f"[QWEN] Модель {model} выдала ошибку: {exc}")
    return []

# ----------------- ПАРСИНГ MARKDOWN -----------------
class HTMLTableExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tables: list[list[list[str]]] = []
        self.current_table: list[list[str]] = []
        self.current_row: list[str] = []
        self.current_cell: list[str] = []
        self.in_cell = False

    def handle_starttag(self, tag, attrs):
        if tag.lower() == "table":
            self.current_table = []
        elif tag.lower() == "tr":
            self.current_row = []
        elif tag.lower() in ("td", "th"):
            self.in_cell = True
            self.current_cell = []

    def handle_endtag(self, tag):
        if tag.lower() in ("td", "th"):
            self.in_cell = False
            self.current_row.append(html.unescape("".join(self.current_cell).strip()))
        elif tag.lower() == "tr":
            if any(c.strip() for c in self.current_row):
                self.current_table.append(self.current_row)
        elif tag.lower() == "table" and self.current_table:
            self.tables.append(self.current_table)

    def handle_data(self, data):
        if self.in_cell:
            self.current_cell.append(data)

def markdown_to_elements(markdown: str) -> list[dict]:
    if not markdown or not markdown.strip():
        return []
    html_pattern = re.compile(r"<table[\s\S]*?</table>", re.IGNORECASE)
    segments = []
    last_idx = 0
    for match in html_pattern.finditer(markdown):
        s, e = match.span()
        if s > last_idx:
            segments.append(("text", markdown[last_idx:s]))
        parser = HTMLTableExtractor()
        parser.feed(match.group(0))
        if parser.tables:
            segments.append(("table", parser.tables[0]))
        last_idx = e
    if last_idx < len(markdown):
        segments.append(("text", markdown[last_idx:]))

    elements: list[dict] = []
    _BOLD_RE = re.compile(r"\*\*(.+?)\*\*|__(.+?)__")

    for seg_type, content in segments:
        if seg_type == "table":
            if content:
                cols = max(len(r) for r in content)
                rows = [r + [""] * (cols - len(r)) for r in content]
                elements.append(make_table(rows=rows, header_row=True, table_kind="other"))
            continue

        for line in content.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
            t = line.strip()
            if not t or t.startswith("```") or re.match(r"^!\[.*?\]\(.*?\)$", t):
                continue
            if re.match(r"^(-{3,}|\*{3,}|_{3,})$", t):
                elements.append(make_paragraph(text="---"))
                continue
            head_m = re.match(r"^(#{1,6})\s+(.*)$", t)
            if head_m:
                lvl = len(head_m.group(1))
                txt = head_m.group(2).strip()
                elements.append(make_heading(text=txt, level=lvl, font_size=max(12.0, 18.0 - (lvl - 1) * 1.5)))
                continue
            runs = []
            pos = 0
            for m in _BOLD_RE.finditer(t):
                if m.start() > pos:
                    runs.append({"text": t[pos:m.start()], "bold": False})
                runs.append({"text": m.group(1) or m.group(2) or "", "bold": True})
                pos = m.end()
            if pos < len(t):
                runs.append({"text": t[pos:], "bold": False})
            clean_txt = "".join(r["text"] for r in runs).strip()
            elements.append(make_paragraph(text=clean_txt, runs=runs if len(runs) > 1 else None))
    return elements

# ----------------- ГЕОМЕТРИЯ СТРАНИЦЫ -----------------
def estimate_margins(page) -> dict:
    rect = page.rect
    w, h = rect.width, rect.height
    min_x, min_y, max_x, max_y = w, h, 0.0, 0.0
    found = False
    for b in (page.get_text("blocks") or []):
        if b[2] > b[0] and b[3] > b[1]:
            found = True
            min_x, min_y = min(min_x, b[0]), min(min_y, b[1])
            max_x, max_y = max(max_x, b[2]), max(max_y, b[3])
    if not found:
        return {"left": DEFAULT_MARGIN_CM, "right": DEFAULT_MARGIN_CM, "top": DEFAULT_MARGIN_CM, "bottom": DEFAULT_MARGIN_CM}
    return {
        "left": min(3.0, max(0.8, round(min_x * PT_TO_CM, 2))),
        "right": min(3.0, max(0.8, round((w - max_x) * PT_TO_CM, 2))),
        "top": min(3.0, max(0.8, round(min_y * PT_TO_CM, 2))),
        "bottom": min(3.0, max(0.8, round((h - max_y) * PT_TO_CM, 2))),
    }

def enrich_elements(page, elements: list[dict]) -> list[dict]:
    page_w = page.rect.width
    for el in elements:
        txt = (el.get("text") or "").strip()
        if not txt or el.get("type") == "table":
            continue
        words = txt.split()[:4]
        match_query = " ".join(words)
        rects = page.search_for(match_query) if match_query else []
        if not rects and words:
            rects = page.search_for(words[0])
        if rects:
            r = rects[0]
            center_dist = abs(((r.x0 + r.x1) / 2.0) - (page_w / 2.0))
            if center_dist < page_w * 0.08 and len(txt) < 90:
                el["align"] = "center"
            elif (page_w - r.x1) < page_w * 0.08 and r.x0 > page_w * 0.40:
                el["align"] = "right"
            else:
                el["align"] = "left"
    return elements

# ----------------- РЕНДЕР DOCX -----------------
def _set_run_font(run, font_name: str, size_pt: float, bold: bool = False):
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.name = font_name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:cs"), font_name)

def render_docx_pages(pages: list[dict]) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = DEFAULT_FONT_NAME
    style.font.size = Pt(9.5)

    for idx, page in enumerate(pages):
        page_w = float(page.get("width_cm") or PAGE_WIDTH_CM)
        page_h = float(page.get("height_cm") or PAGE_HEIGHT_CM)
        orient = page.get("orientation") or ("landscape" if page_w > page_h else "portrait")
        margins = page.get("margins_cm") or {}

        section = doc.sections[0] if idx == 0 else doc.add_section(WD_SECTION_START.NEW_PAGE)
        section.page_width = Cm(page_w)
        section.page_height = Cm(page_h)
        section.orientation = WD_ORIENT.LANDSCAPE if orient == "landscape" else WD_ORIENT.PORTRAIT
        section.left_margin = Cm(float(margins.get("left", DEFAULT_MARGIN_CM)))
        section.right_margin = Cm(float(margins.get("right", DEFAULT_MARGIN_CM)))
        section.top_margin = Cm(float(margins.get("top", DEFAULT_MARGIN_CM)))
        section.bottom_margin = Cm(float(margins.get("bottom", DEFAULT_MARGIN_CM)))

        usable_w = max(10.0, page_w - float(margins.get("left", DEFAULT_MARGIN_CM)) - float(margins.get("right", DEFAULT_MARGIN_CM)))

        elements = page.get("elements", [])
        total_lines = len(elements) + sum(len(el.get("rows", [])) for el in elements if el.get("type") == "table")
        font_scale = 0.82 if total_lines > 42 else (0.88 if total_lines > 32 else 0.95)

        for el in elements:
            etype = el.get("type")
            if etype == "heading":
                p = doc.add_paragraph()
                p.alignment = _ALIGN_MAP.get(el.get("align", "left"), WD_ALIGN_PARAGRAPH.LEFT)
                run = p.add_run(el.get("text", ""))
                _set_run_font(run, DEFAULT_FONT_NAME, max(10.0, (12.0 if el.get("level", 1) <= 1 else 10.5) * font_scale), bold=True)
            elif etype == "table":
                rows = el.get("rows") or []
                if not rows:
                    continue
                cols = max(len(r) for r in rows)
                t = doc.add_table(rows=0, cols=cols)
                t.style = "Table Grid" if el.get("bordered", True) else "Normal Table"

                weights = []
                for ci in range(cols):
                    lens = [len(str(r[ci])) for r in rows if ci < len(r) and r[ci]]
                    weights.append(max(1.5, math.sqrt(sum(lens)/max(len(lens), 1) if lens else 4) * 1.5))
                tot_w = sum(weights) or 1.0
                col_widths = [usable_w * (w / tot_w) for w in weights]

                for r_idx, row in enumerate(rows):
                    cells = t.add_row().cells
                    is_h = bool(el.get("header_row")) and r_idx == 0
                    for c_idx in range(cols):
                        txt = str(row[c_idx] if c_idx < len(row) else "").replace("\n", " ").strip()
                        c = cells[c_idx]
                        c.text = ""
                        p = c.paragraphs[0]
                        p.paragraph_format.space_before = Pt(0)
                        p.paragraph_format.space_after = Pt(0)
                        if re.match(r"^[\d\s,.]+(?:[a-zA-Zа-яА-Я%]+)?$", txt) and not is_h:
                            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        elif len(txt) <= 4 or is_h:
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        else:
                            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        run = p.add_run(txt)
                        _set_run_font(run, DEFAULT_FONT_NAME, max(7.0, 8.5 * font_scale + (0.5 if is_h else 0)), bold=is_h)
                        if c_idx < len(col_widths):
                            c.width = Cm(col_widths[c_idx])
            else:
                txt = el.get("text", "").strip()
                if not txt:
                    continue
                p = doc.add_paragraph()
                p.alignment = _ALIGN_MAP.get(el.get("align", "left"), WD_ALIGN_PARAGRAPH.LEFT)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0 if total_lines > 30 else 0.5)
                runs = el.get("runs")
                size = max(7.0, min(float(el.get("font_size", 9.5)) * font_scale, 13.0))
                if runs:
                    for r in runs:
                        if r.get("text"):
                            run = p.add_run(r["text"])
                            _set_run_font(run, DEFAULT_FONT_NAME, size, bold=bool(r.get("bold", False)))
                else:
                    run = p.add_run(txt)
                    _set_run_font(run, DEFAULT_FONT_NAME, size, bold=bool(el.get("bold", False)))

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()

# ----------------- ТОЧКА ВХОДА ДЛЯ DOCUMENT_ROUTES -----------------
def _prepare_page(pdf_bytes: bytes, page_num: int, on_done) -> dict:
   
    w_cm = h_cm = 0.0
    try:
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")  # свой экземпляр на поток
        try:
            page = doc[page_num - 1]
            rect = page.rect
            w_cm, h_cm = rect.width * PT_TO_CM, rect.height * PT_TO_CM
            margins_cm = estimate_margins(page)
            pix = page.get_pixmap(dpi=PAGE_DPI, alpha=False)
            png_bytes = pix.tobytes("png")
            raw_text = page.get_text("text") or ""
        finally:
            doc.close()

        # 1. LightOn OCR
        ocr_md = ocr_page_lighton(png_bytes)

        # 2. Qwen Refinement (структура)
        elements = []
        if ocr_md:
            elements = refine_page_qwen(ocr_md, png_bytes)

        # 3. Fallbacks
        if not elements and ocr_md:
            elements = markdown_to_elements(ocr_md)
        if not elements:
            elements = [make_paragraph(text=t) for t in raw_text.splitlines() if t.strip()] or [make_paragraph(text="[Текст не найден]")]

        return {
            "page_number": page_num,
            "margins_cm": margins_cm,
            "width_cm": round(w_cm, 2),
            "height_cm": round(h_cm, 2),
            "orientation": "landscape" if w_cm > h_cm else "portrait",
            "elements": elements,
        }
    finally:
        on_done()


def export_modified_pdf_to_docx_vlm(modified_doc: fitz.Document, doc_title: str = "document", on_progress=None) -> io.BytesIO:
    total_pages = modified_doc.page_count


    buf = io.BytesIO()
    modified_doc.save(buf, garbage=3, deflate=True)
    pdf_bytes = buf.getvalue()

    # Потокобезопасный счётчик завершённых страниц (Задача 4, п.7).
    _lock = threading.Lock()
    _counter = {"n": 0}

    def _on_page_done():
        with _lock:
            _counter["n"] += 1
            done = _counter["n"]
        if on_progress:
            pct = int((done / max(1, total_pages)) * 90)
            on_progress(pct, f"Распознавание страниц... ({done}/{total_pages})")


    workers = max(VLM_MAX_WORKERS, DOCX_MAX_WORKERS)

    pages = []
    if total_pages <= 0:
        pass
    elif workers <= 1:
        for page_num in range(1, total_pages + 1):
            pages.append(_prepare_page(pdf_bytes, page_num, _on_page_done))
    else:
        with ThreadPoolExecutor(max_workers=workers) as pool:
            futures = [
                pool.submit(_prepare_page, pdf_bytes, page_num, _on_page_done)
                for page_num in range(1, total_pages + 1)
            ]
            for fut in futures:
                pages.append(fut.result())

    # ВОССТАНОВЛЕНИЕ порядка страниц после параллельной обработки (Задача 4, п.4).
    pages.sort(key=lambda p: p["page_number"])

    if on_progress:
        on_progress(95, "Сборка чистого документа Word...")

   
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    try:
        for p in pages:
            page = doc[p["page_number"] - 1]
            p["elements"] = enrich_elements(page, p["elements"])
    finally:
        doc.close()

    docx_bytes = render_docx_pages(pages)
    stream = io.BytesIO(docx_bytes)
    stream.seek(0)
    return stream