# -*- coding: utf-8 -*-
"""
docx_export.py
Экспорт страниц UDM в .docx с сохранением таблиц, выравнивания и плавающих печатей.
"""

import base64
import io
import logging
import math
import re

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn

logger = logging.getLogger("docx_export")


def _safe_color(hex_str):
    if not hex_str:
        return None
    h = hex_str.lstrip("#").strip()
    if len(h) != 6:
        return None
    try:
        return RGBColor.from_string(h.upper())
    except Exception:
        return None


def _apply_font(run, font_meta, font_scale: float = 1.0):
    if not font_meta:
        return
    size_pt = font_meta.get("size_pt")
    if size_pt:
        try:
            run.font.size = Pt(max(7.0, min(24.0, float(size_pt) * font_scale)))
        except Exception:
            pass
    weight = font_meta.get("weight", 400)
    run.font.bold = bool(weight and weight >= 600)
    color = _safe_color(font_meta.get("color"))
    if color is not None:
        run.font.color.rgb = color
    family = font_meta.get("family", "Arial")
    if family:
        run.font.name = family


def _detect_line_alignment(line_bbox: dict, page_w_pt: float, margin_left_pt: float = 30.0, margin_right_pt: float = 30.0):
    """
    Математический расчет выравнивания строки на основе её геометрии относительно листа.
    """
    if not line_bbox:
        return WD_ALIGN_PARAGRAPH.LEFT, 0.0

    x0 = float(line_bbox.get("x1", 0))
    x1 = float(line_bbox.get("x2", 0))
    line_w = max(1.0, x1 - x0)
    page_mid = page_w_pt / 2.0
    line_mid = (x0 + x1) / 2.0
    printable_w = page_w_pt - margin_left_pt - margin_right_pt

    # 1. По центру: середина строки близка к середине листа (до 9%), и строка не на всю ширину
    if abs(line_mid - page_mid) < (page_w_pt * 0.09) and line_w < (printable_w * 0.85):
        return WD_ALIGN_PARAGRAPH.CENTER, 0.0

    # 2. По правому краю: строка в правой половине листа
    if (page_w_pt - margin_right_pt - x1) < (page_w_pt * 0.10) and x0 > (page_w_pt * 0.45):
        return WD_ALIGN_PARAGRAPH.RIGHT, 0.0

    # 3. Красная строка (абзацный отступ): умеренный отступ (от 12 до 60 pt)
    indent_pt = 0.0
    if 12.0 < (x0 - margin_left_pt) < 60.0:
        indent_pt = min(28.0, (x0 - margin_left_pt) * 0.6)

    return WD_ALIGN_PARAGRAPH.LEFT, indent_pt


def _render_line_spans(p, spans, default_font, font_scale=1.0):
    if not spans:
        return
    for sp in spans:
        sp_text = sp.get("text", "") if isinstance(sp, dict) else getattr(sp, "text", "")
        if not sp_text:
            continue
        run = p.add_run(sp_text)

        fam = (sp.get("font_family") if isinstance(sp, dict) else getattr(sp, "font_family", None)) or default_font.get("family", "Times New Roman")
        size = float((sp.get("size_pt") if isinstance(sp, dict) else getattr(sp, "size_pt", None)) or default_font.get("size_pt", 11.0)) * font_scale
        bold = bool((sp.get("is_bold") if isinstance(sp, dict) else getattr(sp, "is_bold", False)) or (default_font.get("weight", 400) >= 600))
        italic = bool(sp.get("is_italic") if isinstance(sp, dict) else getattr(sp, "is_italic", False))
        color_hex = (sp.get("color") if isinstance(sp, dict) else getattr(sp, "color", None)) or default_font.get("color")

        run.font.name = fam
        run.font.size = Pt(max(7.0, min(24.0, size)))
        run.font.bold = bold
        run.font.italic = italic
        c = _safe_color(color_hex)
        if c:
            run.font.color.rgb = c


def _render_line_item(doc, line_item, page_w_pt: float = 595.32, font_scale: float = 1.0, is_compact: bool = False):
    """Рендерит один элемент: обычный абзац или двухколоночную строку на одном уровне Y."""
    if line_item.get("type") == "split_row":
        left_item = line_item["left"]
        right_item = line_item["right"]

        t = doc.add_table(rows=1, cols=2)
        t.style = "Normal Table"
        _set_cell_margins(t.cell(0, 0), top=0, bottom=0, left=0, right=0)
        _set_cell_margins(t.cell(0, 1), top=0, bottom=0, left=0, right=0)

        p_left = t.cell(0, 0).paragraphs[0]
        p_right = t.cell(0, 1).paragraphs[0]
        p_left.paragraph_format.space_before = Pt(0)
        p_left.paragraph_format.space_after = Pt(0.5 if is_compact else 1.2)
        p_right.paragraph_format.space_before = Pt(0)
        p_right.paragraph_format.space_after = Pt(0.5 if is_compact else 1.2)
        p_left.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        _render_line_spans(p_left, left_item.get("spans") or [{"text": left_item.get("text", "")}], left_item.get("font", {}), font_scale=font_scale)
        _render_line_spans(p_right, right_item.get("spans") or [{"text": right_item.get("text", "")}], right_item.get("font", {}), font_scale=font_scale)
        return

    text = (line_item.get("text") or "").strip()
    if not text:
        return

    bbox = line_item.get("bbox") or {}
    align, indent_pt = _detect_line_alignment(bbox, page_w_pt)

    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0.5 if is_compact else 1.2)
    p.paragraph_format.line_spacing = 1.0 if is_compact else 1.05
    if indent_pt > 0 and align == WD_ALIGN_PARAGRAPH.LEFT:
        p.paragraph_format.first_line_indent = Pt(indent_pt)

    spans = line_item.get("spans") or []
    if spans:
        _render_line_spans(p, spans, line_item.get("font", {}), font_scale=font_scale)
    else:
        run = p.add_run(text)
        _apply_font(run, line_item.get("font", {}), font_scale=font_scale)


def _set_cell_margins(cell, top=30, bottom=30, left=50, right=50):
    try:
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)
    except Exception:
        pass


def _compute_smart_column_widths(rows_data, max_cols, total_width_inches=7.2):
    weights = []
    for ci in range(max_cols):
        lens = []
        for row in rows_data:
            cells = row.get("cells", []) if isinstance(row, dict) else getattr(row, "cells", [])
            if ci < len(cells):
                c_item = cells[ci]
                txt = str(c_item.get("text", "") if isinstance(c_item, dict) else getattr(c_item, "text", "")).strip()
                if txt:
                    lens.append(len(txt))
        avg = sum(lens) / max(len(lens), 1) if lens else 2
        # Нелинейный вес (квадратный корень) дает широким столбцам место, не пережимая узкие
        w = max(1.0, math.sqrt(avg) * 1.6)
        weights.append(w)

    total_weight = sum(weights) or 1.0
    return [(w / total_weight) * total_width_inches for w in weights]


def _adjust_column_widths(table, rows_data, max_cols):
    try:
        col_widths_in = _compute_smart_column_widths(rows_data, max_cols)
        for row in table.rows:
            for c_idx, cell in enumerate(row.cells):
                if c_idx < len(col_widths_in):
                    cell.width = Inches(col_widths_in[c_idx])
    except Exception as e:
        logger.warning(f"[DOCX] adjust widths error: {e}")


def _add_table_block(doc, block, font_scale: float = 1.0, is_compact: bool = False):
    rows_data = block.get("rows") if isinstance(block, dict) else getattr(block, "rows", [])
    if not rows_data:
        raw = (block.get("raw_text") if isinstance(block, dict) else getattr(block, "raw_text", "")) or ""
        if raw.strip():
            for line in raw.split("\n"):
                if line.strip():
                    doc.add_paragraph(line)
        return

    num_rows = len(rows_data)
    max_cols = 0
    for r in rows_data:
        cells = r.get("cells", []) if isinstance(r, dict) else getattr(r, "cells", [])
        if len(cells) > max_cols:
            max_cols = len(cells)

    if max_cols == 0 or num_rows == 0:
        return

    table = doc.add_table(rows=num_rows, cols=max_cols)
    is_bordered = bool(block.get("bordered", True) if isinstance(block, dict) else getattr(block, "bordered", True))
    
    # Установка границ таблицы
    try:
        if is_bordered:
            table.style = "Table Grid"
        else:
            table.style = "Normal Table"
    except Exception:
        pass

    base_cell_size = max(7.0, min(8.5 * font_scale, 9.5))

    for r_idx, row in enumerate(rows_data):
        cells_list = row.get("cells", []) if isinstance(row, dict) else getattr(row, "cells", [])
        for c_idx in range(max_cols):
            cell = table.cell(r_idx, c_idx)
            _set_cell_margins(cell, top=20 if is_compact else 30, bottom=20 if is_compact else 30)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0

            if c_idx < len(cells_list):
                c_item = cells_list[c_idx]
                text_val = str(c_item.get("text", "") if isinstance(c_item, dict) else getattr(c_item, "text", "")).strip()
                is_header = bool(c_item.get("is_header") if isinstance(c_item, dict) else getattr(c_item, "is_header", False)) or (r_idx == 0 and is_bordered and max_cols > 2)
            else:
                text_val = ""
                is_header = False

            # Выравнивание текста
            if re.match(r"^[\d\s,.]+(?:[a-zA-Zа-яА-Я%₽$€]+)?$", text_val) and not is_header and len(text_val) > 2:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif len(text_val) <= 4 or is_header:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            run = p.add_run(text_val)
            run.font.size = Pt(base_cell_size + (0.5 if is_header else 0.0))
            run.font.name = "Arial"
            if is_header:
                run.font.bold = True

    # Расчёт ширин колонок
    dict_rows = []
    for r in rows_data:
        cells = r.get("cells", []) if isinstance(r, dict) else getattr(r, "cells", [])
        dict_rows.append({"cells": [{"text": (c.get("text", "") if isinstance(c, dict) else getattr(c, "text", ""))} for c in cells]})
    _adjust_column_widths(table, dict_rows, max_cols)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(0.5 if is_compact else 1.5)


def _add_floating_picture(paragraph, image_bytes: bytes, width_cm: float, left_cm: float, top_cm: float):
    try:
        run = paragraph.add_run()
        img_stream = io.BytesIO(image_bytes)
        run.add_picture(img_stream, width=Cm(width_cm))
        drawing = run._element.xpath(".//w:drawing")[0]
        inline = drawing.xpath(".//wp:inline")[0]
        extent = inline.xpath(".//wp:extent")[0]
        cx, cy = extent.get("cx"), extent.get("cy")
        docPr = inline.xpath(".//wp:docPr")[0]
        doc_id = docPr.get("id", "1")
        doc_name = docPr.get("name", "Stamp")
        graphic = inline.xpath(".//a:graphic")[0]

        left_emu = int(left_cm * 360000)
        top_emu = int(top_cm * 360000)
        ns_wp = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
        ns_a = "http://schemas.openxmlformats.org/drawingml/2006/main"
        ns_r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

        anchor = parse_xml(f'''
        <wp:anchor xmlns:wp="{ns_wp}" xmlns:a="{ns_a}" xmlns:r="{ns_r}"
            distT="0" distB="0" distL="0" distR="0"
            simplePos="0" relativeHeight="251658240" behindDoc="0" locked="0"
            layoutInCell="0" allowOverlap="1">
          <wp:simplePos x="0" y="0"/>
          <wp:positionH relativeFrom="page"><wp:posOffset>{left_emu}</wp:posOffset></wp:positionH>
          <wp:positionV relativeFrom="page"><wp:posOffset>{top_emu}</wp:posOffset></wp:positionV>
          <wp:extent cx="{cx}" cy="{cy}"/>
          <wp:effectExtent l="0" t="0" r="0" b="0"/>
          <wp:wrapNone/>
          <wp:docPr id="{doc_id}" name="{doc_name}"/>
          <wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>
        </wp:anchor>''')
        anchor.append(graphic)
        drawing.replace(inline, anchor)
    except Exception as e:
        logger.warning(f"[DOCX] Ошибка наложения плавающей картинки: {e}")


def _add_picture_block(doc, block, page_w_pt: float = 595.32, page_h_pt: float = 841.92):
    b64 = block.get("image_base64")
    if not b64 or "rel_x" not in block:
        return

    try:
        if b64.startswith("data:"):
            b64 = b64.split(",", 1)[1]
        img_bytes = base64.b64decode(b64)

        page_w_cm = page_w_pt * 2.54 / 72.0
        page_h_cm = page_h_pt * 2.54 / 72.0
        left_cm = float(block["rel_x"]) * page_w_cm
        top_cm = float(block["rel_y"]) * page_h_cm
        width_cm = float(block.get("rel_w", 0.22)) * page_w_cm

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        _add_floating_picture(p, img_bytes, width_cm=width_cm, left_cm=left_cm, top_cm=top_cm)
    except Exception as exc:
        logger.warning(f"[DOCX] Ошибка вставки печати: {exc}")


def build_docx_from_udm_pages(pages: list) -> io.BytesIO:
    doc = Document()

    for section in doc.sections:
        section.left_margin = Pt(30)
        section.right_margin = Pt(30)
        section.top_margin = Pt(30)
        section.bottom_margin = Pt(30)

    first_page = True
    for page in pages:
        if not first_page:
            doc.add_page_break()
        first_page = False

        blocks = page.get("blocks") or []
        page_w_pt = float(page.get("width_pt", 595.32))
        page_h_pt = float(page.get("height_pt", 841.92))

        # 1. Извлекаем элементы: многострочные блоки сохраняем как цельные абзацы, а не рвём на куски
        flat_elements = []
        for block in blocks:
            btype = block.get("type")
            if btype in ("table", "picture"):
                bbox = block.get("bbox") or {}
                y0 = float(bbox.get("y1", 0) if isinstance(bbox, dict) else getattr(bbox, "y1", 0))
                flat_elements.append({"type": btype, "block": block, "y0": y0, "x0": 0.0})
            else:
                lines = block.get("lines") or []
                default_font = block.get("font") or {}
                if not lines:
                    raw = block.get("raw_text") or ""
                    lines_text = [t for t in raw.split("\n") if t.strip()]
                    lines = [{"id": f"{block.get('id', 'b')}_l{i}", "text": t, "bbox": block.get("bbox")} for i, t in enumerate(lines_text)]

                # Если это многострочный абзац (текст договора) — объединяем строки и спаны с пробелами
                if len(lines) > 1:
                    first_l_bbox = lines[0].get("bbox") or {}
                    merged_spans = []
                    for li_idx, line in enumerate(lines):
                        spans = line.get("spans") or [{"text": line.get("text", "")}]
                        for sp in spans:
                            merged_spans.append(sp)
                        if li_idx < len(lines) - 1:
                            merged_spans.append({"text": " ", "font_family": default_font.get("family"), "size_pt": default_font.get("size_pt")})

                    flat_elements.append({
                        "type": "text_line",
                        "text": " ".join((l.get("text") or "").strip() for l in lines if (l.get("text") or "").strip()),
                        "bbox": first_l_bbox,
                        "spans": merged_spans,
                        "font": default_font,
                        "is_paragraph": True,
                        "y0": float(first_l_bbox.get("y1", 0) if isinstance(first_l_bbox, dict) else getattr(first_l_bbox, "y1", 0)),
                        "x0": float(first_l_bbox.get("x1", 0) if isinstance(first_l_bbox, dict) else getattr(first_l_bbox, "x1", 0))
                    })
                else:
                    for line in lines:
                        l_bbox = line.get("bbox") or {}
                        y0 = float(l_bbox.get("y1", 0) if isinstance(l_bbox, dict) else getattr(l_bbox, "y1", 0))
                        x0 = float(l_bbox.get("x1", 0) if isinstance(l_bbox, dict) else getattr(l_bbox, "x1", 0))
                        flat_elements.append({
                            "type": "text_line",
                            "text": line.get("text", ""),
                            "bbox": l_bbox,
                            "spans": line.get("spans") or [],
                            "font": default_font,
                            "is_paragraph": False,
                            "y0": y0,
                            "x0": x0
                        })

        # 2. Сортируем все элементы сверху вниз по Y, затем слева направо по X
        flat_elements.sort(key=lambda item: (round(item.get("y0", 0), 1), round(item.get("x0", 0), 1)))

        # 3. Кластеризуем строки на одном горизонтальном уровне (разница по Y <= 4 pt и разнос по горизонтали > 100 pt)
        clustered = []
        skip_indices = set()
        for i in range(len(flat_elements)):
            if i in skip_indices:
                continue
            cur = flat_elements[i]

            if cur["type"] == "text_line" and i + 1 < len(flat_elements):
                nxt = flat_elements[i + 1]
                if nxt["type"] == "text_line" and abs(cur["y0"] - nxt["y0"]) <= 4.0:
                    x0_first = cur["x0"]
                    x0_second = nxt["x0"]
                    if abs(x0_first - x0_second) > 100.0:
                        left_el = cur if x0_first < x0_second else nxt
                        right_el = nxt if x0_first < x0_second else cur
                        clustered.append({
                            "type": "split_row",
                            "left": left_el,
                            "right": right_el
                        })
                        skip_indices.add(i + 1)
                        continue

            clustered.append(cur)

        total_lines = len(clustered)
        is_compact = total_lines > 25
        font_scale = 0.85 if total_lines > 38 else (0.92 if total_lines > 28 else 1.0)

        # 4. Рендерим отсортированные и сгруппированные блоки
        for item in clustered:
            itype = item.get("type")
            try:
                if itype == "table":
                    _add_table_block(doc, item["block"], font_scale=font_scale, is_compact=is_compact)
                elif itype == "picture":
                    _add_picture_block(doc, item["block"], page_w_pt=page_w_pt, page_h_pt=page_h_pt)
                else:
                    _render_line_item(doc, item, page_w_pt=page_w_pt, font_scale=font_scale, is_compact=is_compact)
            except Exception as exc:
                logger.exception(f"[DOCX] Ошибка рендеринга ({itype}): {exc}")

    if not pages:
        doc.add_paragraph("")

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out